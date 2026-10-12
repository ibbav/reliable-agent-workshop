#!/usr/bin/env python3
import argparse
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from build_deck import NOTES

ORANGE = RGBColor(254, 86, 7)
BLACK = RGBColor(20, 20, 20)
GRAY = RGBColor(95, 95, 95)

TITLES = [
    "Nothing crashed. It was just quietly wrong.",
    "A built-in devil's advocate",
    "What we version together",
    "How we develop a change",
    "Real cases, variants, and simulated conversations",
    "Who judges, and how certain is the difference?",
    "Make the error visible in the workflow",
    "How reviewed experience changes the next version",
    "Thank you. Questions?"
]

def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)

def style_document(document):
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(13)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size in [("Title", 30), ("Heading 1", 19), ("Heading 2", 11)]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 2" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        if style_name == "Title":
            paragraph_properties = style._element.get_or_add_pPr()
            for border in paragraph_properties.findall(qn("w:pBdr")):
                paragraph_properties.remove(border)
    add_page_number(section.footer.paragraphs[0])

def add_label(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = ORANGE

def add_slide_page(document, number, title, image_path, notes):
    document.add_page_break()
    add_label(document, f"Slide {number} of 9")
    heading = document.add_heading(title, level=1)
    heading.paragraph_format.space_after = Pt(6)
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_after = Pt(6)
    picture.add_run().add_picture(str(image_path), width=Inches(5.65))
    paragraphs = [item.strip() for item in notes.split("\n\n") if item.strip()]
    for index, text in enumerate(paragraphs):
        if text.startswith("Timing:"):
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width = Inches(1.4)
            table.columns[1].width = Inches(4.85)
            first, second = text.split("\n", 1) if "\n" in text else (text, "")
            table.cell(0, 0).text = first.replace("Timing:", "TIMING")
            table.cell(0, 1).text = second.replace("Visual cue:", "VISUAL CUE")
            for cell in table.rows[0].cells:
                set_cell_shading(cell, "F3F3F3")
                for run in cell.paragraphs[0].runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    run.bold = True
            continue
        paragraph = document.add_paragraph(text)
        if text.startswith("Transition:") or text.startswith("Honesty note:") or text.startswith("Sources:"):
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(4)
            for run in paragraph.runs:
                run.font.color.rgb = ORANGE
                run.bold = True

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--slides-dir", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    document = Document()
    style_document(document)
    add_label(document, "Reliable Agent Workshop")
    title = document.add_heading("Working Products 2026\nEnglish speaker script", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = document.add_paragraph("Valentina Ibba · 30 minutes · Nine slides")
    subtitle.runs[0].font.size = Pt(15)
    subtitle.runs[0].font.color.rgb = ORANGE
    document.add_paragraph("This printable script contains the full PowerPoint notes, slide thumbnails, timing, visual cues, transitions, caveats, and sources. All workshop cases and replay results are synthetic.")
    summary = document.add_table(rows=4, cols=2)
    rows = [
        ("Total timing", "30:00"),
        ("Replay math", "8 cases × 3 repeats × 2 versions = 48 offline fixture attempts"),
        ("Candidate fixture result", "24 of 24 contract checks; 21 of 24 expected-behavior checks"),
        ("Boundary", "No live model performance claim; no external action is performed")
    ]
    for row, (label, value) in zip(summary.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "F3F3F3")
        row.cells[0].paragraphs[0].runs[0].bold = True
    document.add_paragraph("How to use this script: read the bold timing and visual cue first, then use the prose as a speakable path. Orange transition and honesty notes mark the claims that deserve extra care.")
    slides_dir = Path(args.slides_dir)
    for index, (title_text, notes) in enumerate(zip(TITLES, NOTES), 1):
        add_slide_page(document, index, title_text, slides_dir / f"slide-{index}.png", notes)
    document.add_page_break()
    add_label(document, "Public sources")
    document.add_heading("Sources and honesty notes", level=1)
    sources = [
        "Anthropic, Building effective agents: https://www.anthropic.com/research/building-effective-agents",
        "Anthropic, Prompt engineering overview: https://docs.anthropic.com/en/docs/build-with-claude/prompt-engineering/overview",
        "Gelman et al., Bayesian Data Analysis, third edition.",
        "Zhang et al., Agentic Context Engineering: https://arxiv.org/abs/2510.04618"
    ]
    for source in sources:
        document.add_paragraph(source, style="List Bullet")
    document.add_heading("Honesty notes", level=2)
    for note in [
        "The conference assistant, sessions, messages, outputs, and reported replay results are synthetic.",
        "The offline replay makes zero network calls and does not measure model variability.",
        "The 67 percent example is a separate teaching illustration using 17 of 20 versus 18 of 20.",
        "Passing finite cases does not prove general reliability. Protected-rule failures remain vetoes.",
        "The repository is local only. No publication or paid model call is part of this build."
    ]:
        document.add_paragraph(note, style="List Bullet")
    document.core_properties.title = "Working Products 2026: English Speaker Script"
    document.core_properties.subject = "Reliable Agent Workshop"
    document.core_properties.author = "Valentina Ibba"
    document.core_properties.last_modified_by = "Valentina Ibba"
    document.core_properties.comments = "Synthetic public workshop edition."
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)

if __name__ == "__main__":
    main()
